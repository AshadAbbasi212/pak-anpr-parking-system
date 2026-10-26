"""Station screen (local web page) + JSON API."""
from __future__ import annotations

import csv
import io
import logging
import re
from datetime import datetime
from pathlib import Path

from flask import Flask, Response, jsonify, render_template, request, send_from_directory

from .plates import PlateFormat, clean, display
from .printing import MODES, PrintQueue, PrinterSettings, list_windows_printers
from .store import Store, StoreError, as_aware, local_day

log = logging.getLogger("web")

CAMERA_LABELS = {"camera1": "Camera 1 (gate)", "camera2": "Camera 2 (road)", "manual": "manual"}


def _local(dt: datetime | None) -> str:
    dt = as_aware(dt)
    return dt.astimezone().strftime("%Y-%m-%d %H:%M:%S") if dt else ""


def entry_json(e: dict) -> dict:
    return {
        "id": e["_id"], "kind": e["kind"], "slip_code": e.get("slip_code") or "", "plate": e["plate"],
        "plate_display": display(e["plate"]), "vehicle_type": e.get("vehicle_type", ""),
        "group": e.get("vehicle_group", ""), "time": _local(e.get("created_at")),
        "camera": CAMERA_LABELS.get(e.get("camera"), e.get("camera") or ""),
        "snapshot": e.get("snapshot"), "plate_img": e.get("plate_img"),
        "print_status": e.get("print_status", ""), "print_count": e.get("print_count", 0),
        "print_error": e.get("print_error", ""), "handed_over": bool(e.get("handed_over")),
        "voided": bool(e.get("voided")), "notes": e.get("notes") or "",
    }


def create_app(cfg: dict, store: Store, printq: PrintQueue, printer_settings: PrinterSettings,
               workers: dict, fmt: PlateFormat) -> Flask:
    app = Flask(__name__, template_folder=str(Path(__file__).parent / "templates"))
    data_dir = Path(cfg["storage"]["data_dir"])
    snap_dir = data_dir / "snapshots"
    slips_dir = data_dir / "slips"

    def bad(msg: str, code: int = 400):
        return jsonify({"error": msg}), code

    def read_plate(raw: str):
        plate = fmt.normalize_manual(raw) if fmt.enabled else clean(raw)
        if not plate:
            return None, "That does not look like a valid plate number (formats are set in config.yaml > validation)"
        return plate, None

    # ------------------------------------------------------------------ pages / media
    @app.get("/")
    def index():
        return render_template("index.html", cameras=list(workers))

    @app.get("/snapshots/<path:rel>")
    def snapshots(rel):
        return send_from_directory(snap_dir, rel)

    @app.get("/slips/<path:name>")
    def slip_png(name):
        return send_from_directory(slips_dir, name)

    @app.get("/slip/<entry_id>")
    def slip_page(entry_id):
        e = store.get_entry(entry_id)
        if not e or e["kind"] != "slip":
            return "Slip not found", 404
        return render_template("slip.html", code=e["slip_code"])

    @app.get("/stream/<cam>")
    def stream(cam):
        worker = workers.get(cam)
        if worker is None:
            return bad("unknown camera", 404)

        def gen():
            seq = 0
            while True:
                jpeg, seq = worker.get_jpeg(seq, 2.0)
                if jpeg is None:
                    continue
                yield (b"--frame\r\nContent-Type: image/jpeg\r\nContent-Length: " + str(len(jpeg)).encode()
                       + b"\r\n\r\n" + jpeg + b"\r\n")

        return Response(gen(), mimetype="multipart/x-mixed-replace; boundary=frame",
                        headers={"Cache-Control": "no-store"})

    # ------------------------------------------------------------------ live status
    @app.get("/api/status")
    def status():
        events = [{
            "time": _local(e.get("time")).split(" ")[-1], "camera": CAMERA_LABELS.get(e["camera"], e["camera"]),
            "plate": e.get("plate", ""), "color": e["color"], "label": e["label"], "reason": e["reason"],
        } for e in list(store.events)[:25]]
        with store._lock:
            unacked = [a for a in store.alerts if not a.get("ack")]
        loud = [a for a in unacked if a.get("severity") == "danger"]
        loud.sort(key=lambda a: a["time"])
        return jsonify({
            "stats": store.stats(), "cameras": [w.status() for w in workers.values()],
            "db_online": store.online, "pending_writes": store.pending_count,
            "events": events, "unacked_alerts": len(unacked), "loud_alerts": len(loud),
            "latest_loud_id": loud[-1]["_id"] if loud else "",
            "printer": printq.status(),
        })

    # ------------------------------------------------------------------ slip queue + actions
    @app.get("/api/queue")
    def queue():
        rows = store.queue(include_handed=request.args.get("handed") == "1")
        return jsonify([entry_json(e) for e in rows])

    @app.post("/api/entries/<entry_id>/handed")
    def handed(entry_id):
        value = bool((request.get_json(silent=True) or {}).get("value", True))
        try:
            store.set_handed_over(entry_id, value)
        except StoreError as exc:
            return bad(str(exc), 404)
        return jsonify({"ok": True})

    @app.post("/api/entries/<entry_id>/void")
    def void(entry_id):
        try:
            store.void_entry(entry_id)
        except StoreError as exc:
            return bad(str(exc), 404)
        return jsonify({"ok": True})

    @app.post("/api/entries/<entry_id>/reprint")
    def reprint(entry_id):
        e = store.get_entry(entry_id)
        if not e or e["kind"] != "slip":
            return bad("Slip not found", 404)
        if e.get("voided"):
            return bad("This slip has been voided", 400)
        store._update(entry_id, {"print_status": "queued", "print_error": ""})
        printq.enqueue(entry_id, reprint=True)
        return jsonify({"ok": True})

    @app.post("/api/entries/<entry_id>/edit_plate")
    def edit_plate(entry_id):
        plate, err = read_plate((request.get_json(silent=True) or {}).get("plate", ""))
        if err:
            return bad(err)
        try:
            store.edit_plate(entry_id, plate)
        except StoreError as exc:
            return bad(str(exc), 404)
        return jsonify({"ok": True})

    # ------------------------------------------------------------------ history / export / summary
    def _filters():
        return (request.args.get("q", ""), request.args.get("kind", ""), request.args.get("from", ""),
                request.args.get("to", ""), request.args.get("voided") == "1")

    @app.get("/api/history")
    def history():
        page = max(1, int(request.args.get("page", 1)))
        per = min(200, max(1, int(request.args.get("per", 25))))
        q, kind, d_from, d_to, voided = _filters()
        rows, total, complete = store.query_entries(q, kind, d_from, d_to, voided, (page - 1) * per, per)
        return jsonify({"rows": [entry_json(e) for e in rows], "total": total, "page": page, "per": per,
                        "complete": complete})

    @app.get("/api/export")
    def export():
        q, kind, d_from, d_to, voided = _filters()
        try:
            rows = store.export_entries(q, kind, d_from, d_to, voided)
        except StoreError as exc:
            return bad(str(exc), 503)
        header = ["Slip no.", "Plate", "Vehicle", "Date and time", "Camera", "Type", "Print status",
                  "Handed over", "Voided", "Notes"]
        data = []
        for e in rows:
            j = entry_json(e)
            data.append([j["slip_code"], j["plate_display"], j["vehicle_type"], j["time"], j["camera"],
                         "whitelisted (no slip)" if j["kind"] == "whitelisted" else "slip", j["print_status"],
                         "yes" if j["handed_over"] else "no", "yes" if j["voided"] else "no", j["notes"]])
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        if request.args.get("format") == "xlsx":
            from openpyxl import Workbook
            wb = Workbook()
            ws = wb.active
            ws.title = "Entries"
            ws.append(header)
            for r in data:
                ws.append(r)
            for col, width in zip("ABCDEFGHIJ", (12, 14, 10, 20, 17, 20, 13, 12, 8, 50)):
                ws.column_dimensions[col].width = width
            buf = io.BytesIO()
            wb.save(buf)
            return Response(buf.getvalue(),
                            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                            headers={"Content-Disposition": f"attachment; filename=entries_{stamp}.xlsx"})
        out = io.StringIO()
        w = csv.writer(out)
        w.writerow(header)
        w.writerows(data)
        return Response("\ufeff" + out.getvalue(), mimetype="text/csv; charset=utf-8",
                        headers={"Content-Disposition": f"attachment; filename=entries_{stamp}.csv"})

    @app.get("/api/summary")
    def summary():
        try:
            return jsonify(store.hourly_summary(request.args.get("date") or local_day()))
        except StoreError as exc:
            return bad(str(exc))

    # ------------------------------------------------------------------ whitelist
    @app.get("/api/whitelist")
    def whitelist_get():
        q = clean(request.args.get("q", ""))
        with store._lock:
            items = [w for w in store.whitelist.values() if q in w["plate"]]
        items.sort(key=lambda w: w.get("added_at") or datetime.min, reverse=True)
        return jsonify([{"plate": w["plate"], "plate_display": display(w["plate"]), "note": w.get("note", ""),
                         "added_at": _local(w.get("added_at"))} for w in items])

    @app.post("/api/whitelist")
    def whitelist_add():
        body = request.get_json(silent=True) or {}
        plate, err = read_plate(body.get("plate", ""))
        if err:
            return bad(err)
        store.whitelist_add(plate, (body.get("note") or "")[:200])
        return jsonify({"ok": True})

    @app.delete("/api/whitelist/<plate>")
    def whitelist_remove(plate):
        try:
            store.whitelist_remove(clean(plate))
        except StoreError as exc:
            return bad(str(exc), 404)
        return jsonify({"ok": True})

    @app.post("/api/whitelist/import")
    def whitelist_import():
        f = request.files.get("file")
        if f is None or not f.filename:
            return bad("Choose a .csv or .xlsx file first")
        name = f.filename.lower()
        rows: list[list] = []
        try:
            if name.endswith(".xlsx"):
                from openpyxl import load_workbook
                ws = load_workbook(io.BytesIO(f.read()), read_only=True, data_only=True).active
                rows = [list(r) for r in ws.iter_rows(values_only=True)]
            elif name.endswith((".csv", ".txt")):
                text = f.read().decode("utf-8-sig", errors="replace")
                rows = list(csv.reader(io.StringIO(text)))
            else:
                return bad("Only .csv or .xlsx files are supported")
        except Exception as exc:  # noqa: BLE001
            return bad(f"Could not read that file: {exc}")
        items, rejected = [], []
        for i, row in enumerate(rows):
            if not row or row[0] is None or str(row[0]).strip() == "":
                continue
            raw = str(row[0]).strip()
            if i == 0 and re.fullmatch(r"(?i)\s*(plates?|number|plate number|vehicle)\s*", raw):
                continue                                       # header row
            plate, err = read_plate(raw)
            note = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ""
            if err:
                rejected.append(raw)
            else:
                items.append((plate, note[:200]))
        n = store.whitelist_add_many(items)
        return jsonify({"ok": True, "added": n, "rejected": rejected[:30], "rejected_count": len(rejected)})

    # ------------------------------------------------------------------ alerts
    @app.get("/api/alerts")
    def alerts_get():
        show_all = request.args.get("all") == "1"
        with store._lock:
            rows = [a for a in store.alerts if show_all or not a.get("ack")]
        rows = sorted(rows, key=lambda a: a["time"], reverse=True)[:150]
        return jsonify([{
            "id": a["_id"], "time": _local(a["time"]), "type": a["type"], "severity": a.get("severity", "warning"),
            "message": a["message"], "plate": a.get("plate"), "camera": CAMERA_LABELS.get(a.get("camera"), ""),
            "snapshot": a.get("snapshot"), "entry_id": a.get("entry_id"), "ack": bool(a.get("ack")),
        } for a in rows])

    @app.post("/api/alerts/<alert_id>/ack")
    def alerts_ack(alert_id):
        try:
            store.ack_alert(alert_id)
        except StoreError as exc:
            return bad(str(exc), 404)
        return jsonify({"ok": True})

    @app.post("/api/alerts/ack_all")
    def alerts_ack_all():
        with store._lock:
            ids = [a["_id"] for a in store.alerts if not a.get("ack")]
        for i in ids:
            try:
                store.ack_alert(i)
            except StoreError:
                pass
        return jsonify({"ok": True, "count": len(ids)})

    # ------------------------------------------------------------------ printer settings
    @app.get("/api/printer")
    def printer_get():
        names, default = list_windows_printers()
        return jsonify({"settings": printer_settings.get(), "printers": names, "default_printer": default,
                        "modes": list(MODES)})

    @app.post("/api/printer")
    def printer_set():
        try:
            return jsonify({"ok": True, "settings": printer_settings.update(request.get_json(silent=True) or {})})
        except StoreError as exc:
            return bad(str(exc))

    @app.post("/api/printer/test")
    def printer_test():
        ok, msg = printq.test_print()
        return jsonify({"ok": ok, "message": msg}), (200 if ok else 500)

    return app
