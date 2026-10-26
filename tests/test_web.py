import io

import pytest
from conftest import snap, ts

from app.plates import PlateFormat
from app.slips import CAM1, CAM2
from app.web import create_app


class FakeWorker:
    def __init__(self, name):
        self.name = name

    def status(self):
        return {"camera": self.name, "status": "online", "fps": 5.0, "last_plate": "", "analysing": True}

    def get_jpeg(self, seq, timeout=1):
        return b"\xff\xd8fake", seq + 1


@pytest.fixture
def client(store, engine, printq, cfg):
    fmt = PlateFormat(True, [r"^[A-Z]{2,4}[0-9]{2,4}$"])
    app = create_app(cfg, store, printq, printq.settings, {CAM1: FakeWorker(CAM1), CAM2: FakeWorker(CAM2)}, fmt)
    return app.test_client()


def test_pages_and_status(client, store, engine):
    assert client.get("/").status_code == 200
    engine.on_plate(CAM1, "LEA1234", 0.9, "CAR", ts(), 100.0, snap)
    j = client.get("/api/status").get_json()
    assert j["stats"]["slips_today"] == 1 and j["db_online"] is True and len(j["cameras"]) == 2
    assert j["events"][0]["plate"] == "LEA-1234" and j["events"][0]["color"] == "green"
    assert client.get("/stream/camera1", buffered=False).status_code == 200
    assert client.get("/stream/nope").status_code == 404


def test_queue_and_slip_actions(client, store, engine):
    engine.on_plate(CAM1, "LEA1234", 0.9, "CAR", ts(), 100.0, snap)
    engine.on_plate(CAM1, "ABC999", 0.9, "BIKE", ts(), 100.0, snap)
    q = client.get("/api/queue").get_json()
    assert [r["slip_code"] for r in q] == ["000002", "000001"]
    first = q[1]["id"]
    assert client.post(f"/api/entries/{first}/handed", json={"value": True}).status_code == 200
    assert len(client.get("/api/queue").get_json()) == 1
    assert len(client.get("/api/queue?handed=1").get_json()) == 2
    assert client.post(f"/api/entries/{first}/edit_plate", json={"plate": "nonsense!"}).status_code == 400
    assert client.post(f"/api/entries/{first}/edit_plate", json={"plate": "lea-1235"}).status_code == 200
    assert client.post(f"/api/entries/{first}/reprint").status_code == 200
    assert client.get(f"/slip/{first}").status_code == 200
    assert client.post(f"/api/entries/{first}/void").status_code == 200
    assert client.post(f"/api/entries/{first}/void").status_code == 404
    assert client.post(f"/api/entries/{first}/reprint").status_code == 400
    assert client.get("/api/status").get_json()["stats"]["slips_today"] == 1


def test_history_export_and_summary(client, store, engine):
    engine.on_plate(CAM1, "LEA1234", 0.9, "CAR", ts(), 100.0, snap)
    engine.on_plate(CAM1, "ABC999", 0.9, "BIKE", ts(), 100.0, snap)
    store.whitelist_add("VIP0001")
    engine.on_plate(CAM1, "VIP0001", 0.9, "CAR", ts(), 100.0, snap)
    h = client.get("/api/history?kind=bike").get_json()
    assert h["total"] == 1 and h["rows"][0]["plate_display"] == "ABC-999" and h["complete"] is True
    assert client.get("/api/history?kind=whitelisted").get_json()["total"] == 1
    assert client.get("/api/history?q=000001").get_json()["total"] == 1
    csv = client.get("/api/export?format=csv")
    assert csv.status_code == 200 and "LEA-1234" in csv.get_data(as_text=True)
    x = client.get("/api/export?format=xlsx")
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(x.data)).active
    assert ws.max_row == 4 and ws["A1"].value == "Slip no."
    s = client.get("/api/summary").get_json()
    assert s["totals"] == {"car": 1, "bike": 1, "other": 0, "whitelisted": 1}
    assert client.get("/api/summary?date=garbage").status_code == 400


def test_whitelist_add_remove_and_import(client, store):
    assert client.post("/api/whitelist", json={"plate": "nonsense!"}).status_code == 400
    assert client.post("/api/whitelist", json={"plate": "lea-1234", "note": "boss"}).status_code == 200
    assert client.get("/api/whitelist").get_json()[0]["plate_display"] == "LEA-1234"
    csv_data = b"Plate,Note\nABC-123,staff\nXYZ 9876,\nthis is bad,\n"
    r = client.post("/api/whitelist/import", data={"file": (io.BytesIO(csv_data), "list.csv")},
                    content_type="multipart/form-data").get_json()
    assert r["added"] == 2 and r["rejected_count"] == 1
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(["plate", "note"])
    ws.append(["DEF-555", "vip"])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    r = client.post("/api/whitelist/import", data={"file": (buf, "list.xlsx")},
                    content_type="multipart/form-data").get_json()
    assert r["added"] == 1
    assert store.is_whitelisted("DEF555") and store.is_whitelisted("ABC123")
    assert client.delete("/api/whitelist/DEF555").status_code == 200
    assert client.delete("/api/whitelist/DEF555").status_code == 404
    assert client.post("/api/whitelist/import", data={}, content_type="multipart/form-data").status_code == 400


def test_alerts_and_loud_flag(client, store, engine):
    engine.on_unreadable(CAM2, "BIKE", ts(), 100.0, snap)
    j = client.get("/api/status").get_json()
    assert j["loud_alerts"] == 1 and j["latest_loud_id"]
    a = client.get("/api/alerts").get_json()
    assert a[0]["severity"] == "danger"
    assert client.post(f"/api/alerts/{a[0]['id']}/ack").status_code == 200
    assert client.get("/api/status").get_json()["loud_alerts"] == 0
    assert client.get("/api/alerts").get_json() == []
    assert len(client.get("/api/alerts?all=1").get_json()) == 1


def test_printer_settings_api(client):
    d = client.get("/api/printer").get_json()
    assert d["settings"]["mode"] == "file" and "network" in d["modes"]
    assert client.post("/api/printer", json={"mode": "network", "network_host": "10.1.1.1"}).status_code == 200
    assert client.post("/api/printer", json={"mode": "bogus"}).status_code == 400
    r = client.post("/api/printer/test")
    assert r.status_code == 500                                  # nothing listens on 10.1.1.1 / not reachable
    client.post("/api/printer", json={"mode": "file"})
    assert client.post("/api/printer/test").status_code == 200
